"""
Resume management for JobApplierBot.
Handles parsing, tailoring, and PDF generation.
"""

from __future__ import annotations

import copy
import logging
import os
import re
import subprocess
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

from PyPDF2 import PdfReader
from fpdf import FPDF

logger = logging.getLogger(__name__)


@dataclass
class ResumeData:
    """Structured resume representation."""
    name: str = ""
    email: str = ""
    phone: str = ""
    summary: str = ""
    experience: List[Dict] = field(default_factory=list)
    education: List[Dict] = field(default_factory=list)
    skills: List[str] = field(default_factory=list)
    projects: List[Dict] = field(default_factory=list)
    certifications: List[str] = field(default_factory=list)
    raw_text: str = ""


class ResumeManager:
    """Manages resume loading, parsing, tailoring, and PDF generation."""

    def __init__(self, resume_path: Optional[str] = None):
        self.resume: Optional[ResumeData] = None
        self.last_tailored: Optional[ResumeData] = None
        self._original_path = resume_path

        if resume_path:
            self.load(resume_path)

    def load(self, path: str) -> ResumeData:
        """Load and parse resume from file (PDF, DOCX, TXT, JSON)."""
        self._original_path = path
        ext = Path(path).suffix.lower()

        if ext == ".pdf":
            self.resume = self._parse_pdf(path)
        elif ext == ".docx":
            self.resume = self._parse_docx(path)
        elif ext == ".json":
            self.resume = self._parse_json(path)
        else:
            self.resume = self._parse_txt(path)

        logger.info(f"Loaded resume: {self.resume.name} ({len(self.resume.raw_text)} chars)")
        return self.resume

    def get_text(self, max_chars: int = 12000) -> str:
        """Get resume text, truncated to max_chars."""
        if not self.resume:
            return ""
        return (self.resume.raw_text or "")[:max_chars]

    # ----- Parsing -----

    def _parse_pdf(self, path: str) -> ResumeData:
        """Extract text from PDF."""
        if not os.path.exists(path):
            logger.warning(f"PDF not found: {path}")
            return ResumeData()

        reader = PdfReader(path)
        parts: List[str] = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
        text = "\n".join(parts).strip()
        return self._parse_text(text)

    def _parse_docx(self, path: str) -> ResumeData:
        """Extract text from DOCX."""
        try:
            from docx import Document
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            return self._parse_text(text)
        except ImportError:
            logger.warning("python-docx not installed. Cannot parse DOCX.")
            return ResumeData()

    def _parse_json(self, path: str) -> ResumeData:
        """Load from JSON resume format."""
        import json
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return ResumeData(
            name=data.get("name", ""),
            email=data.get("contact", {}).get("email", ""),
            phone=data.get("contact", {}).get("phone", ""),
            summary=data.get("summary", ""),
            experience=data.get("experience", []),
            education=data.get("education", []),
            skills=self._flatten_skills(data.get("skills", {})),
            projects=data.get("projects", []),
            certifications=data.get("certifications", []),
            raw_text=str(data),
        )

    def _parse_txt(self, path: str) -> ResumeData:
        """Load from plain text."""
        if not os.path.exists(path):
            return ResumeData()
        with open(path, "r", encoding="utf-8") as f:
            text = f.read()
        return self._parse_text(text)

    def _parse_text(self, text: str) -> ResumeData:
        """Parse unstructured resume text into ResumeData."""
        lines = text.split("\n")
        name = lines[0].strip() if lines else ""

        # Extract email
        email = ""
        emails = re.findall(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text)
        if emails:
            email = emails[0]

        # Extract phone
        phone = ""
        phones = re.findall(r'[\+]?\d[\d\s\-\(\)]{8,}\d', text)
        if phones:
            phone = re.sub(r'[^\d+]', '', phones[0])

        # Extract sections
        sections = self._extract_sections(text)

        return ResumeData(
            name=name,
            email=email,
            phone=phone,
            summary=sections.get("summary", ""),
            experience=sections.get("experience", []),
            education=sections.get("education", []),
            skills=sections.get("skills", []),
            projects=sections.get("projects", []),
            certifications=sections.get("certifications", []),
            raw_text=text,
        )

    def _extract_sections(self, text: str) -> Dict:
        """Extract resume sections using keyword matching."""
        sections = {
            "summary": "",
            "experience": [],
            "education": [],
            "skills": [],
            "projects": [],
            "certifications": [],
        }

        lines = text.split("\n")
        current_section = None

        for line in lines:
            line_lower = line.lower().strip()

            # Skip empty lines
            if not line_lower:
                continue

            # Check if this is a section header (standalone or very short line)
            is_header = len(line_lower.split()) <= 4  # Section headers are typically short

            if is_header and any(keyword == line_lower or (keyword in line_lower and len(line_lower) < 30)
                                 for keyword in ["summary", "objective", "profile"]):
                current_section = "summary"
            elif is_header and any(keyword == line_lower or (keyword in line_lower and len(line_lower) < 30)
                                   for keyword in ["experience", "work history", "employment"]):
                current_section = "experience"
            elif is_header and any(keyword == line_lower or (keyword in line_lower and len(line_lower) < 30)
                                   for keyword in ["education", "academic", "qualification"]):
                current_section = "education"
            elif is_header and any(keyword == line_lower or (keyword in line_lower and len(line_lower) < 30)
                                   for keyword in ["skills", "technical skills", "competencies"]):
                current_section = "skills"
            elif is_header and any(keyword == line_lower or (keyword in line_lower and len(line_lower) < 30)
                                   for keyword in ["projects", "portfolio"]):
                current_section = "projects"
            elif is_header and any(keyword == line_lower or (keyword in line_lower and len(line_lower) < 30)
                                   for keyword in ["certifications", "certificate", "licenses"]):
                current_section = "certifications"
            elif current_section and line.strip():
                if current_section == "summary":
                    sections["summary"] += line + " "
                elif current_section == "skills":
                    skills = re.split(r'[,|;]', line)
                    sections["skills"].extend([s.strip() for s in skills if s.strip()])

        return sections

    @staticmethod
    def _flatten_skills(skills_data) -> List[str]:
        """Flatten a skills dict or list into a flat list of strings."""
        if isinstance(skills_data, list):
            return skills_data
        if isinstance(skills_data, dict):
            flat = []
            for val in skills_data.values():
                if isinstance(val, list):
                    flat.extend(val)
                elif isinstance(val, str):
                    flat.append(val)
            return flat
        return []

    # ----- Tailoring -----

    def tailor_for_jd(self, job_description: str, missing_keywords: List[str],
                       suggestions: List[str] = None) -> ResumeData:
        """Create a tailored copy of the resume with added keywords and suggestions."""
        if not self.resume:
            raise ValueError("No resume loaded")

        tailored = copy.deepcopy(self.resume)

        # Add missing keywords to skills
        existing_lc = {s.strip().lower() for s in (tailored.skills or []) if s}
        for kw in missing_keywords[:25]:
            if kw.lower() not in existing_lc:
                tailored.skills.append(kw)
                existing_lc.add(kw.lower())

        # Enhance summary with first suggestion
        if suggestions:
            for suggestion in suggestions:
                if isinstance(suggestion, str) and suggestion.strip():
                    base = (tailored.summary or "").strip()
                    tailored.summary = (base + " " if base else "") + suggestion.strip()[:300]
                    break

        # Rebuild raw text
        tailored.raw_text = self._build_raw_text(tailored)
        self.last_tailored = tailored
        return tailored

    def _build_raw_text(self, resume: ResumeData) -> str:
        """Build flat text representation for ATS scoring."""
        parts: List[str] = []
        parts.append(resume.name or "")
        parts.append(resume.email or "")
        parts.append(resume.phone or "")
        parts.append(f"Summary: {resume.summary or ''}".strip())

        if resume.skills:
            parts.append(f"Skills: {', '.join(resume.skills)}")

        for exp in resume.experience or []:
            if isinstance(exp, dict):
                parts.append(" ".join([
                    exp.get("title", ""),
                    exp.get("company", ""),
                    exp.get("location", ""),
                    exp.get("description", ""),
                    " ".join(exp.get("responsibilities", [])),
                ]))

        for edu in resume.education or []:
            if isinstance(edu, dict):
                parts.append(" ".join([
                    edu.get("degree", ""),
                    edu.get("institution", ""),
                    edu.get("cgpa", ""),
                ]))

        for proj in resume.projects or []:
            if isinstance(proj, dict):
                parts.append(f"{proj.get('name', '')} {proj.get('description', '')}")

        for cert in resume.certifications or []:
            parts.append(str(cert))

        return "\n".join([p for p in parts if p]).strip()

    # ----- PDF Generation -----

    def generate_pdf(self, resume: Optional[ResumeData] = None,
                     output_path: str = "resume_tailored.pdf") -> str:
        """Generate a PDF from resume data. Falls back to simple PDF if LaTeX unavailable."""
        resume_obj = resume or self.resume
        if not resume_obj:
            raise ValueError("No resume data to generate PDF from")

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)

        # Try LaTeX first
        tex_path = str(Path(output_path).with_suffix(".tex"))
        if self._try_latex(resume_obj, tex_path, output_path):
            return output_path

        # Fallback: simple PDF
        self._generate_simple_pdf(resume_obj, output_path)
        return output_path

    def _generate_simple_pdf(self, resume: ResumeData, output_path: str) -> None:
        """Generate a clean PDF using fpdf2."""
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)

        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()

        epw = pdf.epw

        # Header
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(epw, 10, resume.name or "Resume", new_x="LMARGIN", new_y="NEXT", align="C")

        # Contact
        pdf.set_font("Helvetica", "", 10)
        contact_parts = [p for p in [resume.email, resume.phone] if p]
        if contact_parts:
            pdf.cell(epw, 6, " | ".join(contact_parts), new_x="LMARGIN", new_y="NEXT", align="C")
        pdf.ln(5)

        # Summary
        if resume.summary:
            pdf.set_font("Helvetica", "B", 12)
            pdf.cell(epw, 8, "Summary", new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", "", 10)
            pdf.x = pdf.l_margin
            pdf.multi_cell(epw, 5, resume.summary.strip(), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(3)

        # Skills
        if resume.skills:
            pdf.set_font("Helvetica", "B", 12)
            pdf.cell(epw, 8, "Skills", new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", "", 10)
            pdf.x = pdf.l_margin
            pdf.multi_cell(epw, 5, ", ".join(resume.skills), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(3)

        # Experience
        if resume.experience:
            pdf.set_font("Helvetica", "B", 12)
            pdf.cell(epw, 8, "Experience", new_x="LMARGIN", new_y="NEXT")
            for exp in resume.experience:
                if isinstance(exp, dict):
                    pdf.set_font("Helvetica", "B", 10)
                    title_line = f"{exp.get('title', '')} at {exp.get('company', '')}"
                    pdf.cell(epw, 6, title_line, new_x="LMARGIN", new_y="NEXT")
                    pdf.set_font("Helvetica", "I", 9)
                    dates = f"{exp.get('startDate', '')} - {exp.get('endDate', '')}"
                    pdf.cell(epw, 5, f"{dates} | {exp.get('location', '')}", new_x="LMARGIN", new_y="NEXT")
                    pdf.set_font("Helvetica", "", 10)
                    for resp in exp.get("responsibilities", []):
                        pdf.x = pdf.l_margin
                        pdf.multi_cell(epw, 5, f"• {resp}", new_x="LMARGIN", new_y="NEXT")
                    pdf.ln(2)

        # Education
        if resume.education:
            pdf.set_font("Helvetica", "B", 12)
            pdf.cell(epw, 8, "Education", new_x="LMARGIN", new_y="NEXT")
            for edu in resume.education:
                if isinstance(edu, dict):
                    pdf.set_font("Helvetica", "B", 10)
                    pdf.cell(epw, 6, edu.get("degree", ""), new_x="LMARGIN", new_y="NEXT")
                    pdf.set_font("Helvetica", "", 10)
                    pdf.cell(epw, 5, f"{edu.get('institution', '')} ({edu.get('startDate', '')}-{edu.get('endDate', '')})", new_x="LMARGIN", new_y="NEXT")
                    pdf.ln(2)

        # Projects
        if resume.projects:
            pdf.set_font("Helvetica", "B", 12)
            pdf.cell(epw, 8, "Projects", new_x="LMARGIN", new_y="NEXT")
            for proj in resume.projects:
                if isinstance(proj, dict):
                    pdf.set_font("Helvetica", "B", 10)
                    pdf.cell(epw, 6, proj.get("name", ""), new_x="LMARGIN", new_y="NEXT")
                    pdf.set_font("Helvetica", "", 10)
                    pdf.x = pdf.l_margin
                    pdf.multi_cell(epw, 5, proj.get("description", ""), new_x="LMARGIN", new_y="NEXT")
                    pdf.ln(2)

        pdf.output(output_path)
        logger.info(f"Generated PDF: {output_path}")

    def _try_latex(self, resume: ResumeData, tex_path: str, pdf_path: str) -> bool:
        """Try to compile resume via LaTeX. Returns True if successful."""
        try:
            template = self._build_latex(resume)
            with open(tex_path, "w", encoding="utf-8") as f:
                f.write(template)

            result = subprocess.run(
                ["pdflatex", "-interaction=nonstopmode", Path(tex_path).name],
                cwd=str(Path(tex_path).parent),
                check=True,
                capture_output=True,
            )
            if os.path.exists(pdf_path):
                logger.info("LaTeX resume compiled successfully")
                return True
        except (subprocess.CalledProcessError, FileNotFoundError):
            pass
        return False

    def _build_latex(self, resume: ResumeData) -> str:
        """Build LaTeX source for the resume."""
        esc = self._latex_escape

        exp_entries = ""
        for e in resume.experience or []:
            if isinstance(e, dict):
                exp_entries += r"\cventry{%s}{%s}{%s}{%s}{}{%s}" % (
                    esc(e.get("dates", "")),
                    esc(e.get("title", "")),
                    esc(e.get("company", "")),
                    esc(e.get("location", "")),
                    esc(e.get("description", "")),
                ) + "\n"

        edu_entries = ""
        for e in resume.education or []:
            if isinstance(e, dict):
                edu_entries += r"\cventry{%s}{%s}{%s}{}{%s}{}" % (
                    esc(e.get("endDate", e.get("startDate", ""))),
                    esc(e.get("degree", "")),
                    esc(e.get("institution", "")),
                    esc(e.get("cgpa", "")),
                ) + "\n"

        skills_str = ", ".join(esc(s) for s in (resume.skills or []) if s)

        proj_entries = ""
        for p in resume.projects or []:
            if isinstance(p, dict):
                proj_entries += r"\cventry{%s}{%s}{}{}{}{%s}" % (
                    esc(p.get("year", "")),
                    esc(p.get("name", "")),
                    esc(p.get("description", "")),
                ) + "\n"

        cert_entries = ""
        for c in resume.certifications or []:
            c_name = c if isinstance(c, str) else c.get("name", "")
            cert_entries += r"\cventry{}{%s}{}{}{}{}" % esc(c_name) + "\n"

        return r"""
\documentclass[11pt,a4paper]{moderncv}
\moderncvstyle{classic}
\moderncvcolor{blue}
\usepackage[scale=0.75]{geometry}

\name{%s}{}
\phone[mobile]{%s}
\email{%s}

\begin{document}
\makecvtitle

\section{Summary}
%s

\section{Experience}
%s

\section{Education}
%s

\section{Skills}
%s

\section{Projects}
%s

\section{Certifications}
%s

\end{document}
""" % (
            esc(resume.name), esc(resume.phone), esc(resume.email),
            esc(resume.summary or ""),
            exp_entries, edu_entries, skills_str, proj_entries, cert_entries,
        )

    @staticmethod
    def _latex_escape(text) -> str:
        """Escape LaTeX special characters."""
        if text is None:
            return ""
        s = str(text)
        return (
            s.replace("\\", r"\textbackslash{}")
            .replace("&", r"\&")
            .replace("%", r"\%")
            .replace("#", r"\#")
            .replace("_", r"\_")
            .replace("{", r"\{")
            .replace("}", r"\}")
            .replace("~", r"\textasciitilde{}")
            .replace("^", r"\^{}")
        )
